import openpyxl, sys
SRC = r"D:/Downloads/PAJAMAS_TOILET_SEAT (1).xlsm"
PID = sys.argv[1] if len(sys.argv) > 1 else '12563'
XLS = r"C:/Users/Administrator/.openclaw/workspace-dev/skills/hicustom-api/output/" + PID + "/listing/listing_filled.xlsm"

# 1) Data Definitions → field -> required status
wb = openpyxl.load_workbook(SRC, data_only=True)
dd = wb["Data Definitions"]
req = {}
label = {}
for r in range(3, dd.max_row + 1):
    f = dd.cell(row=r, column=2).value  # Field name
    required = dd.cell(row=r, column=6).value  # Required?
    lab = dd.cell(row=r, column=3).value
    if f:
        req[str(f).strip()] = str(required or "").strip()
        label[str(f).strip()] = str(lab or "").strip()

# 2) filled template row: col -> field ; row7 values
wb2 = openpyxl.load_workbook(XLS, keep_vba=True)
ws = wb2["Template"]
fields = {}
vals = {}
for c in range(1, ws.max_column + 1):
    f = ws.cell(row=5, column=c).value
    if f:
        fields[c] = str(f).strip()
        vals[str(f).strip()] = ws.cell(row=7, column=c).value

print("=== 必填/条件必填 字段核对（" + PID + "）===")
filled_req, missing_req, missing_cond = [], [], []
for f, status in req.items():
    if status in ("Required", "Conditionally required"):
        v = vals.get(f)
        filled = v is not None and str(v).strip() != ""
        rec = "已填" if filled else "❌ 空"
        if filled:
            filled_req.append(f)
        elif status == "Required":
            missing_req.append((f, label.get(f, f)))
        else:
            missing_cond.append((f, label.get(f, f)))
print("\n[已填必填] 数量", len(filled_req))
print("[缺失·必填 Required] 数量", len(missing_req), "->", [l for f, l in missing_req])
print("[缺失·条件必填 Conditionally required] 数量", len(missing_cond), "->", [l for f, l in missing_cond])
